from io import BytesIO
from pathlib import Path
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches, Pt
from .common import DEFAULT_LOGO, sanitize_text

def format_docx(text, doc_type, terms=None, logo_path=None, company_name="LegalEase"):
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.7); section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8); section.right_margin = Inches(0.8)
    doc.styles["Normal"].font.name = "Times New Roman"
    doc.styles["Normal"].font.size = Pt(11)

    chosen_logo = Path(logo_path) if logo_path else DEFAULT_LOGO
    if chosen_logo.exists() and chosen_logo.suffix.lower() in {".png", ".jpg", ".jpeg"}:
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.add_run().add_picture(str(chosen_logo), width=Inches(1.1))

    title = doc.add_paragraph(); title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run(sanitize_text(doc_type).upper()); run.bold = True; run.font.name = "Times New Roman"; run.font.size = Pt(16)

    for raw in sanitize_text(text).splitlines():
        line = raw.strip()
        if not line: continue
        p = doc.add_paragraph()
        if line.startswith("#"):
            r = p.add_run(line.lstrip("# ").strip()); r.bold = True; r.font.size = Pt(13)
        elif line.startswith(("1.", "2.", "3.", "4.", "5.", "6.", "7.", "8.", "9.")):
            p.add_run(line).bold = True
        else:
            p.add_run(line)

    if terms:
        doc.add_paragraph().add_run("Key Terms").bold = True
        table = doc.add_table(rows=1, cols=2); table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.style = "Table Grid"
        table.rows[0].cells[0].text = "No."; table.rows[0].cells[1].text = "Term"
        for idx, term in enumerate(terms, 1):
            cells = table.add_row().cells; cells[0].text = str(idx); cells[1].text = sanitize_text(term)

    footer = section.footer.paragraphs[0]; footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run(f"{company_name} - AI-generated draft. Review before use.")
    buf = BytesIO(); doc.save(buf); return buf.getvalue()
